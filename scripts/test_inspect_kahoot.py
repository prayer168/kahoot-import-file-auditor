"""Behavior checks using in-memory workbook objects, without rewriting user files."""
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook
from inspect_kahoot import check_item, inspect_bytes, inspect_workbook, iter_inputs, load_rules


HEADERS = [None, 'Question - max 95 characters', 'Answer 1', 'Answer 2',
           'Answer 3', 'Answer 4', 'Time limit (sec)', 'Correct answer(s)']
RULES = dict(source_url='https://example.test/verified-test-profile', checked_at='2026-09-18',
             question_max=95, answer_max=60, time_values=[5, 10, 20, 30, 60, 120])


def item(**overrides):
    values = dict(question='哪個是證據？', answer1='觀察紀錄', answer2='未驗證猜想',
                  answer3=None, answer4=None, time=20, correct=1)
    values.update(overrides)
    return dict(values=values)


def codes(record):
    return {p['code'] for p in record['issues']}


class InspectorTests(unittest.TestCase):
    def test_two_options_and_numeric_answer_are_valid(self):
        self.assertEqual(check_item(item(), RULES)['issues'], [])
        self.assertEqual(check_item(item(correct=1.0), RULES)['parsed_correct_indices'], [1])

    def test_multiselect_and_invalid_separators(self):
        self.assertEqual(check_item(item(correct='1, 2'), RULES)['parsed_correct_indices'], [1, 2])
        for value in ('1，2', '1;2', '1,5', '', '1.5', True):
            self.assertIn('invalid_answer_indices', codes(check_item(item(correct=value))))

    def test_answer_points_to_blank(self):
        self.assertIn('answer_points_to_blank', codes(check_item(item(correct=4))))

    def test_duplicate_options_and_duplicate_indices(self):
        result = check_item(item(answer2='觀察紀錄', correct='1,1'))
        self.assertTrue({'duplicate_options', 'duplicate_answer_index'} <= codes(result))

    def test_missing_question_and_option_gap(self):
        result = check_item(item(question='', answer2=None, answer4='另一選項'))
        self.assertTrue({'empty_question', 'option_gap'} <= codes(result))

    def test_rules_are_optional_and_explicit(self):
        self.assertNotIn('invalid_time', codes(check_item(item(time=90))))
        self.assertIn('invalid_time', codes(check_item(item(time=90), RULES)))
        self.assertIn('length_limit', codes(check_item(item(question='題' * 96), RULES)))

    def test_formula_is_pending_not_evaluated(self):
        record = item(correct='=1+1')
        record['formula_fields'] = ['correct']
        self.assertEqual(codes(check_item(record, RULES)), {'formula_cell'})

    def test_header_offset_original_coordinates_and_hidden_sheet(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '自然題庫'
        sheet.append(['說明'])
        sheet.append(HEADERS)
        sheet.append([7, '哪個是證據？', '觀察紀錄', '未驗證猜想', None, None, 20, 1])
        sheet.append([8])  # Numbered template placeholder.
        sheet.append([9, None, '甲', '乙', None, None, 20, 1])
        notes = workbook.create_sheet('補充')
        notes.sheet_state = 'hidden'
        notes.append(['範圍待確認'])
        report = inspect_workbook(workbook, 'fixture.xlsx')
        self.assertEqual(len(report['items']), 2)
        first = report['items'][0]
        self.assertEqual(first['excel_row'], 3)
        self.assertEqual(first['cells']['question'], 'B3')
        self.assertEqual(first['extra_cells']['A3'], 7)
        self.assertEqual(report['sheets'][1]['state'], 'hidden')
        self.assertIn('unmapped_sheet', codes(report))
        self.assertIn('empty_question', codes(report['items'][1]))
        self.assertEqual(report['content_review'], 'not_started')

    def test_missing_columns_not_silent(self):
        workbook = Workbook()
        workbook.active.append(['Question', 'Answer 1', 'Answer 2'])
        workbook.active.append(['問題', '甲', '乙'])
        report = inspect_workbook(workbook, 'partial.xlsx')
        self.assertIn('missing_columns', codes(report))
        self.assertNotIn('invalid_answer_indices', codes(report['items'][0]))
        self.assertEqual(report['items'][0]['review_status'], '尚未審查')

    def test_bad_workbook_and_empty_workbook_not_passed(self):
        self.assertEqual(inspect_bytes(b'broken', 'bad.xlsx')['status'], 'read_error')
        report = inspect_workbook(Workbook(), 'empty.xlsx')
        self.assertTrue({'no_items', 'unmapped_sheet'} <= codes(report))

    def test_zip_members_retained_without_extracting(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            archive = folder / 'batch.zip'
            with zipfile.ZipFile(archive, 'w') as z:
                z.writestr('../outside.xlsx', b'not-a-workbook')
                z.writestr('part/two.xlsx', b'also-not-a-workbook')
            records = list(iter_inputs([archive, archive]))
            self.assertEqual(len(records), 2)
            self.assertIn('!../outside.xlsx@', records[0][0])
            self.assertEqual(list(folder.iterdir()), [archive])

    def test_rule_profile_validation(self):
        with tempfile.TemporaryDirectory() as temp:
            profile = Path(temp) / 'rules.json'
            profile.write_text(json.dumps(RULES), encoding='utf-8')
            self.assertEqual(load_rules(profile)['question_max'], 95)
            profile.write_text('{}', encoding='utf-8')
            with self.assertRaises(ValueError):
                load_rules(profile)


if __name__ == '__main__':
    unittest.main()
